import re
import streamlit as st
import anthropic
from docx import Document
from io import BytesIO
from dotenv import load_dotenv
from law_search import TOOLS, run_tool

load_dotenv()

st.set_page_config(page_title="律师法律助手", page_icon="⚖️", layout="wide")
st.title("⚖️ 律师法律助手")
st.caption("上传或粘贴客户沟通记录，自动分析案情并生成法律文书（法条自动从国家法律数据库获取）")

# ── Session Stat初始化 ──────────────────────────────────────────────────────
for key in ("analysis_result", "doc_result", "doc_type_used"):
    if key not in st.session_state:
        st.session_state[key] = None


# ── 文件解析 ──────────────────────────────────────────────────────────────────
def extract_text_from_file(uploaded_file):
    name = uploaded_file.name.lower()
    if name.endswith(".txt"):
        raw = uploaded_file.read()
        for enc in ("utf-8", "gbk", "utf-16"):
            try:
                return raw.decode(enc)
            except UnicodeDecodeError:
                continue
        return raw.decode("utf-8", errors="replace")
    if name.endswith(".docx"):
        doc = Document(uploaded_file)
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    return None


# ── Markdown → Word 转换 ──────────────────────────────────────────────────────
def _add_inline_bold(paragraph, text: str):
    """将含 **bold** 标记的文本以正确格式写入段落。"""
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        else:
            paragraph.add_run(part)


def _markdown_to_docx(doc: Document, text: str):
    """将 Markdown 格式的分析报告逐行转换为 Word 格式。"""
    for line in text.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=2)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=1)
        elif stripped.startswith("# "):
            doc.add_heading(stripped[2:], level=0)
        elif stripped.startswith(("- ", "* ")):
            p = doc.add_paragraph(style="List Bullet")
            _add_inline_bold(p, stripped[2:])
        elif re.match(r"^\d+\.\s", stripped):
            p = doc.add_paragraph(style="List Number")
            _add_inline_bold(p, re.sub(r"^\d+\.\s", "", stripped))
        else:
            p = doc.add_paragraph()
            _add_inline_bold(p, stripped)


def build_word_doc(source_text: str, content: str, title: str) -> BytesIO:
    doc = Document()
    doc.add_heading(title, 0)
    doc.add_heading("案情原文", level=1)
    doc.add_paragraph(source_text)
    doc.add_heading("分析内容", level=1)
    _markdown_to_docx(doc, content)
    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


# ── Claude Agentic Call ───────────────────────────────────────────────────────
def _agentic_call(system: str, user: str, status_container=None) -> str:
    """通用 agentic loop：Claude 自动调用工具直至完成任务。"""
    client = anthropic.Anthropic()
    messages = [{"role": "user", "content": user}]

    while True:
        response = client.messages.create(
            model="claude-sonnet-4-6",
            max_tokens=4096,
            system=system,
            tools=TOOLS,
            messages=messages,
        )

        tool_calls = [b for b in response.content if b.type == "tool_use"]
        if tool_calls and status_container:
            names = "、".join(b.name for b in tool_calls)
            status_container.info(f"正在查询法律数据库：{names}…")

        if response.stop_reason == "end_turn" or not tool_calls:
            return "\n".join(
                b.text for b in response.content if hasattr(b, "text")
            )

        messages.append({"role": "assistant", "content": response.content})
        tool_results = [
            {
                "type": "tool_result",
                "tool_use_id": b.id,
                "content": run_tool(b.name, b.input),
            }
            for b in tool_calls
        ]
        messages.append({"role": "user", "content": tool_results})


def run_analysis(text: str, status_container=None) -> str:
    system = (
        "你是一位资深律师助理，拥有访问中华人民共和国国家法律数据库的能力。"
        "分析案件时，必须主动使用工具查询真实法条，不得凭记忆引用法律条文。"
        "引用法律时需注明具体条款编号。"
    )
    user = f"""请对以下客户咨询内容进行专业分析：

{text}

请按以下结构输出分析报告：

## 一、案件基本事实梳理
（提炼关键事实：时间线、当事人、核心争议点）

## 二、涉及法律领域
（判断案件类型，如劳动纠纷、合同纠纷、婚姻家事等）

## 三、相关法律法规
（使用工具查询真实法条后，列出适用的具体条款）

## 四、裁判文书网检索关键词
（列出5-8个适合在中国裁判文书网搜索的关键词组合）

## 五、类似案件裁判观点总结
（总结近5年同类案件的主流裁判观点和胜诉/败诉规律）

## 六、给客户的初步建议
（用通俗易懂的语言，给出具体可操作的建议）"""
    return _agentic_call(system, user, status_container)


DOC_PROMPTS = {
    "答辩状": (
        "你是一位资深诉讼律师，擅长起草答辩状。"
        "必须使用工具查询真实法条，所有法律引用需注明条款编号。"
        "答辩状语言正式、逻辑严密，符合中国司法文书规范。",
        """请根据以下案情和补充信息，起草一份专业的答辩状。

{text}

{extra}

答辩状须包含以下部分：
1. 答辩人及原告基本信息
2. 答辩请求
3. 事实与理由（分条论述，每条引用具体法律条款）
4. 此致（法院名称）
5. 答辩人签名及日期""",
    ),
    "法律意见书": (
        "你是一位资深法律顾问，擅长撰写法律意见书。"
        "必须使用工具查询真实法条，所有法律引用需注明条款编号。"
        "意见书语言专业、客观，分析有据，结论明确。",
        """请根据以下案情和补充信息，起草一份专业的法律意见书。

{text}

{extra}

法律意见书须包含以下部分：
## 一、委托事项
## 二、基本事实
## 三、法律分析
（逐项分析法律关系，引用具体法条）
## 四、法律意见
（明确结论和建议）
## 五、声明
（本意见书仅供参考，不构成诉讼承诺）
## 六、出具机构及日期""",
    ),
}


def generate_document(
    text: str, doc_type: str, extra_info: str, status_container=None
) -> str:
    system, user_template = DOC_PROMPTS[doc_type]
    extra = f"补充信息：\n{extra_info}" if extra_info.strip() else ""
    user = user_template.format(text=text, extra=extra)
    return _agentic_call(system, user, status_container)


# ── UI：输入区 ────────────────────────────────────────────────────────────────
tab_upload, tab_paste = st.tabs(["📁 上传文件", "📋 手动粘贴"])

input_text = ""

with tab_upload:
    uploaded_file = st.file_uploader(
        "支持微信导出的 .txt 文件，或 .docx 文档",
        type=["txt", "docx"],
    )
    if uploaded_file:
        extracted = extract_text_from_file(uploaded_file)
        if extracted:
            st.success(f"已读取：{uploaded_file.name}（共 {len(extracted)} 字符）")
            with st.expander("预览（前 500 字）"):
                st.text(extracted[:500] + ("…" if len(extracted) > 500 else ""))
            input_text = extracted
        else:
            st.error("无法读取该文件格式，请上传 .txt 或 .docx 文件。")

with tab_paste:
    pasted = st.text_area(
        "请粘贴客户的咨询内容：",
        height=200,
        placeholder="例如：客户说他被公司无故辞退，没有得到任何赔偿…",
    )
    if pasted.strip():
        input_text = pasted

st.divider()

# ── UI：功能区 ────────────────────────────────────────────────────────────────
action_analysis, action_doc = st.tabs(["📋 案情分析", "📄 法律文书生成"])

# ── Tab 1：案情分析 ───────────────────────────────────────────────────────────
with action_analysis:
    if st.button("开始分析", type="primary", key="btn_analysis"):
        if not input_text.strip():
            st.warning("请先上传文件或粘贴咨询内容。")
        else:
            with st.spinner("正在分析中，请稍候…"):
                try:
                    status = st.empty()
                    st.session_state.analysis_result = run_analysis(
                        input_text, status
                    )
                    status.empty()
                except anthropic.AuthenticationError:
                    st.error("API Key 无效，请检查 .env 文件中的 ANTHROPIC_API_KEY。")
                except anthropic.RateLimitError:
                    st.error("API 调用频率超限，请稍后重试。")
                except Exception as e:
                    st.error(f"分析失败：{e}")

    if st.session_state.analysis_result:
        st.subheader("分析结果")
        st.markdown(st.session_state.analysis_result)
        word_buf = build_word_doc(
            input_text, st.session_state.analysis_result, "法律咨询分析报告"
        )
        st.download_button(
            "📄 下载 Word 分析报告",
            data=word_buf,
            file_name="法律咨询分析报告.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

# ── Tab 2：法律文书生成 ───────────────────────────────────────────────────────
with action_doc:
    doc_type = st.selectbox("文书类型", list(DOC_PROMPTS.keys()))
    extra_info = st.text_area(
        "补充信息（可选）",
        height=120,
        placeholder="如：法院名称、当事人全名、案号、需重点论述的争议焦点…",
    )

    if st.button("生成文书", type="primary", key="btn_doc"):
        if not input_text.strip():
            st.warning("请先上传文件或粘贴案情内容。")
        else:
            with st.spinner(f"正在生成{doc_type}，请稍候…"):
                try:
                    status = st.empty()
                    st.session_state.doc_result = generate_document(
                        input_text, doc_type, extra_info, status
                    )
                    st.session_state.doc_type_used = doc_type
                    status.empty()
                except anthropic.AuthenticationError:
                    st.error("API Key 无效，请检查 .env 文件中的 ANTHROPIC_API_KEY。")
                except anthropic.RateLimitError:
                    st.error("API 调用频率超限，请稍后重试。")
                except Exception as e:
                    st.error(f"生成失败：{e}")

    if st.session_state.doc_result:
        st.subheader(f"{st.session_state.doc_type_used} 草稿")
        st.markdown(st.session_state.doc_result)
        word_buf = build_word_doc(
            input_text,
            st.session_state.doc_result,
            st.session_state.doc_type_used,
        )
        st.download_button(
            f"📄 下载 {st.session_state.doc_type_used} Word 文件",
            data=word_buf,
            file_name=f"{st.session_state.doc_type_used}.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
